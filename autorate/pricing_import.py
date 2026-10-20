"""pricing 엑셀 업로드: 파싱, 검수, 기존 자료와 비교, 저장.

기획서 4-1(비용 항목 선택), 4-4(USD만), 7장 5·16·17번(기간 중복, 파일 오류, 재업로드) 기준.
"""

from __future__ import annotations

import datetime as dt
import sqlite3
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import BinaryIO

import openpyxl

from autorate import places

SHEET_NAME = "pricing"

# 내부 필드명 -> 엑셀 헤더
HEADERS = {
    "carrier": "선사",
    "pol": "POL",
    "pod": "POD",
    "container": "컨테이너",
    "charge": "항목",
    "charge_type": "구분",
    "unit": "단위",
    "currency": "통화",
    "cost": "Cost",
    "valid_from": "유효시작",
    "valid_to": "유효종료",
    "remark": "비고",
}
REQUIRED_FIELDS = [f for f in HEADERS if f != "remark"]

CONTAINERS = ("20GP", "40GP", "40HC")
# 항목별 구분·단위 (기획서 6-1). 구분·단위가 맞지 않으면 오류 처리
CHARGE_RULES = {
    "Ocean Freight": ("Freight", "Per Container"),
    "BAF": ("Freight", "Per Container"),
    "THC": ("Local", "Per Container"),
    "DOC Fee": ("Local", "Per B/L"),
    "Seal Fee": ("Local", "Per Container"),
}


@dataclass
class PricingRow:
    row_no: int  # 엑셀 행 번호 (헤더가 1행)
    carrier: str = ""
    pol: str = ""
    pod: str = ""
    container: str = ""
    charge: str = ""
    charge_type: str = ""
    unit: str = ""
    currency: str = ""
    cost: Decimal | None = None
    valid_from: dt.date | None = None
    valid_to: dt.date | None = None
    remark: str = ""
    errors: list[str] = field(default_factory=list)

    @property
    def is_valid(self) -> bool:
        return not self.errors

    @property
    def group(self) -> tuple:
        """같은 선사·구간·컨테이너·항목 (기간 중복 검사 단위)."""
        return (self.carrier, self.pol, self.pod, self.container, self.charge)

    @property
    def key(self) -> tuple:
        """DB 저장 키. 같은 키면 업데이트, 다르면 별개 행."""
        return (*self.group, self.valid_from)


@dataclass
class ParseResult:
    rows: list[PricingRow]
    file_errors: list[str]


@dataclass
class ImportPlan:
    parsed: ParseResult
    new_rows: list[PricingRow]
    changed: list[tuple[PricingRow, dict]]  # (새 행, {"cost": 이전값, "valid_to": 이전값})
    unchanged: list[PricingRow]
    existing_ids: dict  # key -> 기존 rate_id

    @property
    def rows(self) -> list[PricingRow]:
        return self.parsed.rows

    @property
    def file_errors(self) -> list[str]:
        return self.parsed.file_errors

    @property
    def invalid_rows(self) -> list[PricingRow]:
        return [r for r in self.rows if not r.is_valid]


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _parse_date(value) -> dt.date | None:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, str) and value.strip():
        try:
            return dt.date.fromisoformat(value.strip())
        except ValueError:
            return None
    return None


def _parse_decimal(value) -> Decimal | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    if isinstance(value, str):
        try:
            return Decimal(value.replace(",", "").strip())
        except InvalidOperation:
            return None
    return None


def parse_pricing(source: BinaryIO | str) -> ParseResult:
    """pricing 시트를 읽어 행 단위로 검수한다. 파일 자체 오류는 file_errors에 담는다."""
    try:
        wb = openpyxl.load_workbook(source, data_only=True, read_only=True)
    except Exception as exc:  # 엑셀이 아니거나 깨진 파일
        return ParseResult([], [f"엑셀 파일을 열 수 없어요: {exc}"])

    if SHEET_NAME not in wb.sheetnames:
        return ParseResult([], [f"'{SHEET_NAME}' 시트가 없어요. 시트 이름을 확인해 주세요."])

    all_rows = list(wb[SHEET_NAME].iter_rows(values_only=True))
    if not all_rows:
        return ParseResult([], ["pricing 시트가 비어 있어요."])

    header = [_text(h) for h in all_rows[0]]
    col = {}
    missing = []
    for field_name, header_name in HEADERS.items():
        if header_name in header:
            col[field_name] = header.index(header_name)
        elif field_name in REQUIRED_FIELDS:
            missing.append(header_name)
    if missing:
        return ParseResult([], [f"필수 열이 없어요: {', '.join(missing)}. 엑셀 첫 행의 열 이름을 확인해 주세요."])

    def cell(values, name):
        idx = col.get(name)
        return values[idx] if idx is not None and idx < len(values) else None

    rows: list[PricingRow] = []
    for row_no, values in enumerate(all_rows[1:], start=2):
        if all(v is None or _text(v) == "" for v in values):
            continue  # 빈 줄

        row = PricingRow(
            row_no=row_no,
            carrier=_text(cell(values, "carrier")),
            pol=_text(cell(values, "pol")),
            pod=_text(cell(values, "pod")),
            container=_text(cell(values, "container")),
            charge=_text(cell(values, "charge")),
            charge_type=_text(cell(values, "charge_type")),
            unit=_text(cell(values, "unit")),
            currency=_text(cell(values, "currency")),
            cost=_parse_decimal(cell(values, "cost")),
            valid_from=_parse_date(cell(values, "valid_from")),
            valid_to=_parse_date(cell(values, "valid_to")),
            remark=_text(cell(values, "remark")),
        )
        row.errors = _check_row(row)
        rows.append(row)

    _check_duplicates_and_overlaps(rows)
    return ParseResult(rows, [])


def _check_row(row: PricingRow) -> list[str]:
    errors = []
    for name in ("carrier", "pol", "pod", "container", "charge", "unit"):
        if not getattr(row, name):
            errors.append(f"{HEADERS[name]} 값을 입력해 주세요")

    if row.currency != "USD":
        errors.append(f"통화는 USD만 가능해요 (현재 {row.currency or '빈값'})")
    if row.container and row.container not in CONTAINERS:
        errors.append(f"컨테이너는 20GP, 40GP, 40HC 중 하나여야 해요 (현재 {row.container})")

    if row.charge and row.charge not in CHARGE_RULES:
        errors.append(f"항목 이름을 확인해 주세요 (현재 {row.charge})")
    elif row.charge:
        expected_type, expected_unit = CHARGE_RULES[row.charge]
        if row.charge_type != expected_type:
            errors.append(f"{row.charge}의 구분은 {expected_type}이어야 해요")
        if row.unit != expected_unit:
            errors.append(f"{row.charge}의 단위는 {expected_unit}이어야 해요")

    if row.cost is None:
        errors.append("Cost는 숫자로 입력해 주세요")
    elif row.cost < 0:
        errors.append("Cost는 0 이상이어야 해요")

    if row.valid_from is None:
        errors.append("유효시작 날짜를 확인해 주세요")
    if row.valid_to is None:
        errors.append("유효종료 날짜를 확인해 주세요")
    if row.valid_from and row.valid_to and row.valid_to < row.valid_from:
        errors.append("유효종료가 유효시작보다 빨라요. 기간 순서를 확인해 주세요")
    return errors


def _periods_overlap(a: PricingRow, b: PricingRow) -> bool:
    return a.valid_from <= b.valid_to and b.valid_from <= a.valid_to


def _check_duplicates_and_overlaps(rows: list[PricingRow]) -> None:
    """파일 안에서 같은 항목의 기간이 겹치면 겹치는 행 모두 오류로 표시한다 (7장 5번)."""
    valid = [r for r in rows if r.is_valid]
    for i, a in enumerate(valid):
        for b in valid[i + 1 :]:
            if a.group == b.group and _periods_overlap(a, b):
                label = "같은 키 중복" if a.key == b.key else "기간이 겹쳐요"
                a.errors.append(f"{label} (행 {b.row_no})")
                b.errors.append(f"{label} (행 {a.row_no})")


def _stored_rows(conn: sqlite3.Connection) -> list[sqlite3.Row]:
    return conn.execute(
        "SELECT rate_id, carrier, pol, pod, container, charge, cost, valid_from, valid_to "
        "FROM pricing"
    ).fetchall()


def build_plan(parsed: ParseResult, conn: sqlite3.Connection) -> ImportPlan:
    """저장 전 미리보기 계획을 만든다. 기존 저장분과 비교해 신규/변경/동일을 나눈다."""
    stored = _stored_rows(conn)
    existing_ids = {}
    stored_by_key = {}
    for s in stored:
        key = (s["carrier"], s["pol"], s["pod"], s["container"], s["charge"],
               dt.date.fromisoformat(s["valid_from"]))
        existing_ids[key] = s["rate_id"]
        stored_by_key[key] = s

    # 저장된 요율과의 기간 중복 검사 (같은 키는 업데이트이므로 제외)
    for row in parsed.rows:
        if not row.is_valid:
            continue
        for s in stored:
            if (s["carrier"], s["pol"], s["pod"], s["container"], s["charge"]) != row.group:
                continue
            s_from = dt.date.fromisoformat(s["valid_from"])
            s_to = dt.date.fromisoformat(s["valid_to"])
            if s_from == row.valid_from:
                continue
            if row.valid_from <= s_to and s_from <= row.valid_to:
                row.errors.append(
                    f"저장된 요율({s['rate_id']}, {s_from}~{s_to})과 기간이 겹쳐요"
                )
                break

    new_rows, changed, unchanged = [], [], []
    for row in parsed.rows:
        if not row.is_valid:
            continue
        s = stored_by_key.get(row.key)
        if s is None:
            new_rows.append(row)
            continue
        old_cost = Decimal(s["cost"])
        old_to = dt.date.fromisoformat(s["valid_to"])
        if old_cost != row.cost or old_to != row.valid_to:
            changed.append((row, {"cost": old_cost, "valid_to": old_to}))
        else:
            unchanged.append(row)

    return ImportPlan(parsed, new_rows, changed, unchanged, existing_ids)


def unregistered_places(rows: list[PricingRow]) -> list[tuple[int, str]]:
    """지명 변환표에 없는 POL/POD를 (행 번호, 지명)으로 돌려준다. 저장은 막지 않고 알림만 한다."""
    found = []
    for r in rows:
        for name in (r.pol, r.pod):
            if name and not places.is_registered(name):
                found.append((r.row_no, name))
    return found


def status_of(row: PricingRow, as_of: dt.date) -> str:
    """화면 표시용 상태: 오류 / 만료 / 예정 / 정상."""
    if not row.is_valid:
        return "오류"
    if row.valid_to < as_of:
        return "만료"
    if row.valid_from > as_of:
        return "예정"
    return "정상"


def save_plan(
    conn: sqlite3.Connection,
    plan: ImportPlan,
    source_file: str,
    uploaded_at: str,
    skip_invalid: bool = False,
) -> dict:
    """정상 행을 신규 추가/기존 갱신한다. 오류 행이 있으면 skip_invalid=True여야 저장된다."""
    if plan.file_errors:
        raise ValueError("파일 오류가 있어 저장할 수 없습니다.")
    if plan.invalid_rows and not skip_invalid:
        raise ValueError("오류 행이 있습니다. 오류 행을 제외하고 저장하도록 선택하세요.")

    next_no = 1 + max(
        (int(rid[1:]) for rid in plan.existing_ids.values() if rid[1:].isdigit()),
        default=0,
    )
    inserted = updated = 0
    with conn:
        for row in plan.new_rows:
            if not row.is_valid:
                continue
            conn.execute(
                "INSERT INTO pricing (rate_id, carrier, pol, pod, container, charge, charge_type, "
                "unit, currency, cost, valid_from, valid_to, remark, source_file, uploaded_at) "
                "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (
                    f"R{next_no:04d}", row.carrier, row.pol, row.pod, row.container,
                    row.charge, row.charge_type, row.unit, row.currency, str(row.cost),
                    row.valid_from.isoformat(), row.valid_to.isoformat(),
                    row.remark or None, source_file, uploaded_at,
                ),
            )
            next_no += 1
            inserted += 1

        for row, _old in plan.changed:
            conn.execute(
                "UPDATE pricing SET cost = ?, valid_to = ?, remark = ?, source_file = ?, "
                "uploaded_at = ? WHERE rate_id = ?",
                (
                    str(row.cost), row.valid_to.isoformat(), row.remark or None,
                    source_file, uploaded_at, plan.existing_ids[row.key],
                ),
            )
            updated += 1

    return {
        "inserted": inserted,
        "updated": updated,
        "unchanged": len(plan.unchanged),
        "skipped_invalid": len(plan.invalid_rows),
    }


def load_stored(conn: sqlite3.Connection) -> list[dict]:
    rows = conn.execute("SELECT * FROM pricing ORDER BY pol, pod, container, charge, valid_from").fetchall()
    return [dict(r) for r in rows]
